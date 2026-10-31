"""Parsing of tabular data files (CSV/TSV/JSON/YAML/XLSX) for mass printing."""

import csv
import io
import json
import re
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Any, Literal

import yaml
from openpyxl import load_workbook

from app.zpl.templating import VariableDef, merge_variable_defs, validate_value

DataFormat = Literal["csv", "tsv", "json", "yaml", "xlsx"]
COPIES_COLUMNS = ("copies", "_copies")
MAX_COPIES = 1000

_EXTENSIONS: dict[str, DataFormat] = {
    "csv": "csv",
    "txt": "csv",
    "tsv": "tsv",
    "tab": "tsv",
    "json": "json",
    "yaml": "yaml",
    "yml": "yaml",
    "xlsx": "xlsx",
    "xlsm": "xlsx",
}


class DataSourceError(ValueError):
    pass


@dataclass(slots=True)
class DataTable:
    format: DataFormat
    columns: list[str]
    rows: list[dict[str, str]]


@dataclass(slots=True)
class RowResult:
    index: int
    values: dict[str, str]
    copies: int = 1
    errors: list[str] = field(default_factory=list)

    @property
    def valid(self) -> bool:
        return not self.errors


def detect_format(content: bytes, filename: str | None = None) -> DataFormat:
    if filename and "." in filename:
        ext = filename.rsplit(".", 1)[1].lower()
        if ext in _EXTENSIONS:
            return _EXTENSIONS[ext]
    if content[:2] == b"PK":
        return "xlsx"
    head = content.lstrip(b"\xef\xbb\xbf \t\r\n")[:1]
    if head in (b"[", b"{"):
        return "json"
    first_line = content.lstrip(b"\xef\xbb\xbf").split(b"\n", 1)[0]
    if first_line.startswith((b"---", b"- ")):
        return "yaml"
    if b"\t" in first_line and b"," not in first_line:
        return "tsv"
    return "csv"


def _to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, dict | list):
        return json.dumps(value, ensure_ascii=False)
    return str(value).strip()


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def _records_to_table(fmt: DataFormat, records: Iterable[Any]) -> DataTable:
    columns: dict[str, None] = {}
    rows = []
    for position, record in enumerate(records, start=1):
        if not isinstance(record, Mapping):
            raise DataSourceError(f"Row {position} is not an object/mapping")
        row = {str(key).strip(): _to_str(value) for key, value in record.items()}
        for key in row:
            columns.setdefault(key)
        rows.append(row)
    return DataTable(fmt, list(columns), rows)


def _unwrap(data: Any, fmt: str) -> list[Any]:
    if isinstance(data, Mapping) and isinstance(data.get("rows"), list):
        return data["rows"]
    if isinstance(data, list):
        return data
    raise DataSourceError(f"{fmt.upper()} must be a list of objects or {{'rows': [...]}}")


def _parse_delimited(text: str, fmt: DataFormat) -> DataTable:
    if fmt == "tsv":
        delimiter = "\t"
    else:
        sample = text[:8192]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    try:
        header = next(reader)
    except StopIteration:
        raise DataSourceError("The file is empty") from None
    columns = [name.strip() for name in header]
    if not any(columns):
        raise DataSourceError("The header row is empty")
    rows = []
    for record in reader:
        if not any(cell.strip() for cell in record):
            continue
        rows.append(
            {
                column: (record[i].strip() if i < len(record) else "")
                for i, column in enumerate(columns)
                if column
            }
        )
    return DataTable(fmt, [c for c in columns if c], rows)


def _parse_xlsx(content: bytes) -> DataTable:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise DataSourceError(f"Cannot read the Excel file: {exc}") from exc
    try:
        sheet = workbook.worksheets[0]
        values = sheet.iter_rows(values_only=True)
        header = next(values, None)
        if not header or not any(cell is not None for cell in header):
            raise DataSourceError("The first sheet has no header row")
        columns = [_to_str(cell) for cell in header]
        rows = []
        for record in values:
            if not any(cell not in (None, "") for cell in record):
                continue
            rows.append(
                {
                    column: _to_str(record[i]) if i < len(record) else ""
                    for i, column in enumerate(columns)
                    if column
                }
            )
    finally:
        workbook.close()
    return DataTable("xlsx", [c for c in columns if c], rows)


def parse_data(
    content: bytes,
    *,
    filename: str | None = None,
    fmt: DataFormat | None = None,
    max_rows: int = 10_000,
) -> DataTable:
    fmt = fmt or detect_format(content, filename)
    if fmt == "xlsx":
        table = _parse_xlsx(content)
    else:
        text = _decode(content)
        if fmt in ("csv", "tsv"):
            table = _parse_delimited(text, fmt)
        elif fmt == "json":
            try:
                data = json.loads(text)
            except json.JSONDecodeError as exc:
                raise DataSourceError(f"Invalid JSON: {exc}") from exc
            table = _records_to_table("json", _unwrap(data, "json"))
        else:
            try:
                data = yaml.safe_load(text)
            except yaml.YAMLError as exc:
                raise DataSourceError(f"Invalid YAML: {exc}") from exc
            table = _records_to_table("yaml", _unwrap(data, "yaml"))
    if not table.rows:
        raise DataSourceError("The file contains no data rows")
    if len(table.rows) > max_rows:
        raise DataSourceError(f"Too many rows ({len(table.rows)}); the limit is {max_rows}")
    return table


def _normalize(name: str) -> str:
    return re.sub(r"[\s_\-.]+", "", name).lower()


def auto_map(columns: Iterable[str], variables: Iterable[str]) -> dict[str, str | None]:
    """Map each variable to a column with the same (normalized) name."""
    by_norm = {}
    for column in columns:
        by_norm.setdefault(_normalize(column), column)
    return {name: by_norm.get(_normalize(name)) for name in variables}


def copies_column(columns: Iterable[str]) -> str | None:
    for column in columns:
        if column.strip().lower() in COPIES_COLUMNS:
            return column
    return None


def build_rows(
    rows: list[dict[str, str]],
    definitions: list[VariableDef | dict],
    mapping: Mapping[str, str | None],
    constants: Mapping[str, str] | None = None,
    *,
    copies_from: str | None = None,
    field_only: Iterable[str] = (),
) -> list[RowResult]:
    """Resolve each data row into template values and validate it.

    Values of the `field_only` variables (used only in field data) may contain ^ and ~.
    """
    field_only = set(field_only)
    constants = constants or {}
    names = [d.name if isinstance(d, VariableDef) else d["name"] for d in definitions]
    defs = merge_variable_defs(names, definitions)
    results = []
    for index, row in enumerate(rows):
        values: dict[str, str] = {}
        errors: list[str] = []
        for definition in defs:
            column = mapping.get(definition.name)
            if constants.get(definition.name):
                value = constants[definition.name]
            elif column:
                value = row.get(column, "")
            else:
                value = ""
            value = value or definition.default
            error = validate_value(definition, value, allow_control=definition.name in field_only)
            if error:
                errors.append(f"{definition.name} {error}")
            values[definition.name] = value
        copies = 1
        if copies_from and row.get(copies_from, "").strip():
            raw = row[copies_from].strip()
            try:
                number = float(raw)
                if not number.is_integer():
                    raise ValueError
                copies = int(number)
                if not 1 <= copies <= MAX_COPIES:
                    raise ValueError
            except ValueError, OverflowError:
                errors.append(f"copies must be a whole number between 1 and {MAX_COPIES}")
                copies = 1
        results.append(RowResult(index, values, copies, errors))
    return results
